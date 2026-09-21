#!/usr/bin/env python3
"""Build the Git Dojo deck as SVG/HTML and editable PowerPoint shapes.
Requires python-pptx for PowerPoint; all artwork is drawn here, no external assets.
"""
from pathlib import Path
import html
import json
import math

OUT = Path(__file__).resolve().parent
W, H = 1600, 900
C = dict(ink='#172C36', paper='#F5F3EA', white='#FFFFFF', lime='#DAF975', cyan='#65D9EC', coral='#FF8B73', purple='#C5B5F2', muted='#AAB9BD', soft='#E4E8E2', navy='#10232D', panel='#213944')
slides = []

class Slide:
    def __init__(self, title, chapter, notes, dark=False, subtitle=None, minutes=None):
        self.title, self.notes, self.dark, self.minutes = title, notes, dark, minutes
        self.bg = C['navy'] if dark else C['paper']
        self.fg = C['paper'] if dark else C['ink']
        self.shapes = []
        self.rect(0,0,W,H,self.bg,r=0)
        self.text(72,40,'GIT WORKFLOW DOJO',19,C['lime'] if dark else C['ink'],True,w=380)
        self.text(1528,40,chapter.upper(),18,C['muted'] if dark else '#576B70',True,anchor='end',w=700)
        self.text(72,101,title,62 if len(title)<41 else 53,self.fg,True,w=1460)
        if subtitle: self.text(75,185,subtitle,27,C['muted'] if dark else '#576B70',w=1440)
        self.line([(72,835),(1528,835)],'#42555D' if dark else '#D2D9D2',1.5)
        self.text(72,853,'CAMPUS FESTIVAL  /  LERNEN DURCH AUSPROBIEREN',16,C['muted'] if dark else '#576B70',w=1200)
        self.text(1528,849,f'{len(slides)+1:02}',24,self.fg,True,anchor='end',w=90)
        slides.append(self)
    def rect(self,x,y,w,h,fill,stroke=None,r=22,sw=2):
        self.shapes.append(dict(k='rect',x=x,y=y,w=w,h=h,fill=fill,stroke=stroke,r=r,sw=sw))
    def circle(self,x,y,r,fill,stroke=None,sw=3):
        self.shapes.append(dict(k='circle',x=x,y=y,r=r,fill=fill,stroke=stroke,sw=sw))
    def line(self,points,color=None,width=5,dash=False):
        self.shapes.append(dict(k='line',points=points,color=color or self.fg,width=width,dash=dash))
    def text(self,x,y,text,size=30,color=None,bold=False,w=1200,anchor='start',mono=False):
        self.shapes.append(dict(k='text',x=x,y=y,text=text,size=size,color=color or self.fg,bold=bold,w=w,anchor=anchor,mono=mono))
    def pill(self,x,y,label,fill=None,w=None,size=23):
        w=w or max(90,len(label)*size*.61+38)
        self.rect(x,y,w,48,fill or C['lime'],r=24)
        self.text(x+w/2,y+10,label,size,C['ink'],True,w=w-20,anchor='middle',mono=True)
    def node(self,x,y,label,color=None,r=30):
        self.circle(x,y,r,color or C['lime'])
        self.text(x,y-r*.62,label,r*1.07,C['ink'],True,w=r*2,anchor='middle',mono=True)
    def arrow(self,x1,y1,x2,y2,color=None,width=5):
        color=color or self.fg
        self.line([(x1,y1),(x2,y2)],color,width)
        ang=math.atan2(y2-y1,x2-x1)
        pts=[(x2-16*math.cos(ang-.5),y2-16*math.sin(ang-.5)),(x2,y2),(x2-16*math.cos(ang+.5),y2-16*math.sin(ang+.5))]
        self.line(pts,color,width)
    def callout(self,text,fill=None,y=746):
        self.rect(72,y,1456,62,fill or (C['panel'] if self.dark else C['soft']),r=14)
        self.text(99,y+15,text,27,self.fg,w=1385)
    def tag(self,x,y,label,color,down=True,w=None):
        w=w or max(105,len(label)*15+34)
        self.pill(x-w/2,y,label,color,w=w)
        self.line([(x,y+48),(x,y+73)],color,3) if down else self.line([(x,y),(x,y-25)],color,3)

# 01 — hook
s=Slide('Git wird greifbar.', 'Start / 00–05', 'Begrüsse die Gruppe: Wir betreiben gemeinsam die Infoseiten eines Campus-Festivals. Die Inhalte sind klein; wir üben den Weg von einer Idee bis zu einer geprüften Änderung. Zehn Personen, fünf Zweierteams, zwei Stunden. Diese Präsentation wird in kurzen Impulsen zwischen den Praxisphasen genutzt.',True)
s.text(74,260,'Änderungen verstehen.\nGemeinsam weiterbauen.',53,C['paper'],True,w=720)
s.pill(78,438,'10 Personen',C['cyan'],w=220)
s.pill(314,438,'5 Teams',C['purple'],w=180)
s.pill(510,438,'120 Minuten',C['lime'],w=240)
s.line([(850,586),(1020,586),(1150,430),(1340,430),(1470,586)],C['cyan'],12)
s.line([(850,586),(1470,586)],C['lime'],12)
s.line([(1020,586),(1150,704),(1340,704),(1470,586)],C['purple'],12)
for x,y,c in [(850,586,'lime'),(1020,586,'lime'),(1150,430,'cyan'),(1340,430,'cyan'),(1150,704,'purple'),(1340,704,'purple'),(1470,586,'lime')]: s.circle(x,y,24,C[c],C['navy'],6)
s.text(78,631,'BRANCH  →  PR  →  REVIEW  →  MERGE',24,C['muted'],True,mono=True,w=700)
s.text(78,690,'Ein Dojo für die Zusammenarbeit.',31,C['paper'],w=700)

# 02 — motivation
s=Slide('Welche Version gilt?', 'Prinzip 01 / History', 'Zeige die linke Seite und frage: Wer kennt final_final? Git macht Änderungen, ihre Reihenfolge und ihren Kontext nachvollziehbar. Ein Commit ist ein Projektstand mit Metadaten und Elternverweisen, keine weitere Datei namens final. Der Diff hilft beim Vergleich.',subtitle='Aus Dateichaos wird eine nachvollziehbare Geschichte.')
for i,(name,col) in enumerate([('programm_final.md','cyan'),('programm_final2.md','purple'),('programm_WIRKLICH.md','coral')]):
    s.rect(90+i*45,295+i*109,575,122,C[col],r=16)
    s.text(118+i*45,333+i*109,name,30,C['ink'],True,mono=True,w=530)
s.arrow(756,483,865,483,C['ink'],6)
s.line([(966,500),(1420,500)],C['ink'],7)
for x,l in [(980,'A'),(1200,'B'),(1420,'C')]: s.node(x,500,l,C['lime'],37)
s.text(980,571,'Start',28,w=180,anchor='middle')
s.text(1200,571,'Ort ergänzt',28,w=200,anchor='middle')
s.text(1420,571,'Zeit präzisiert',28,w=220,anchor='middle')
s.pill(1080,350,'nachvollziehbare Schritte',C['soft'],w=425,size=23)
s.callout('Jeder Schritt hat einen Stand, eine Begründung und einen Platz in der History.')

# 03 — distributed
s=Slide('Git ist lokal. GitHub verbindet.', 'Prinzip 02 / Lokal & Remote', 'Git ist das Versionskontrollsystem und läuft auch ohne Netzwerk. GitHub hostet Repositories und ergänzt PRs, Reviews und Rechte. Push überträgt Objekte und aktualisiert den Server-Branch. Fetch holt Objekte und aktualisiert Remote-Tracking-Refs; er integriert nicht den aktuellen lokalen Branch. Frage: Ist ein Commit schon auf GitHub? Nein.',subtitle='Ein Commit funktioniert auch ohne Internet.')
s.rect(92,278,570,367,C['ink'],r=24)
s.rect(112,298,530,296,C['white'],r=12)
s.pill(141,326,'DEIN LAPTOP',C['soft'],w=235)
s.text(142,407,'Dateien + Git-History',36,C['ink'],True,w=480)
s.line([(158,539),(570,539)],C['ink'],5)
for x,l in [(175,'A'),(353,'B'),(555,'C')]: s.node(x,539,l,C['lime'],23)
s.rect(67,651,620,27,C['ink'],r=10)
s.rect(960,278,550,400,C['panel'],r=32)
s.pill(1000,315,'GITHUB',C['cyan'],w=165)
s.text(1000,409,'Gemeinsames Repo\nPull Requests\nReviews',38,C['paper'],True,w=450)
s.arrow(699,402,923,402,C['ink'],6)
s.text(810,348,'push',28,C['ink'],True,w=200,anchor='middle',mono=True)
s.arrow(923,545,699,545,C['ink'],6)
s.text(810,578,'fetch',28,C['ink'],True,w=200,anchor='middle',mono=True)
s.callout('Commit ≠ Push. Ein PR ist der gemeinsame Raum für Prüfung und Integration.')

# 04 — staging workflow
s=Slide('Ändern. Auswählen. Festhalten.', 'Prinzip 03 / Die drei Bereiche', 'Working Tree: Dateien bearbeiten. Staging Area oder Index: den nächsten Commit bewusst vorbereiten. Repository: Commit dauerhaft in der lokalen History halten. Die Farben sind nur Lernhilfen, die Beschriftungen tragen die Bedeutung. Mit git diff die ungestagten, mit git diff --staged die gestagten Änderungen prüfen. Neu bearbeitete Inhalte müssen nach dem Staging erneut gestagt werden.',subtitle='Drei Bereiche auf deinem Rechner. Drei verschiedene Zustände.')
for x,n,title,sub,col in [(72,'01','Working Tree','bearbeitete Dateien','cyan'),(580,'02','Staging Area','Auswahl für den Commit','purple'),(1088,'03','Repository','festgehaltener Projektstand','lime')]:
    s.rect(x,287,440,360,C['white'],r=26)
    s.pill(x+24,310,n,C[col],w=72,size=23)
    s.text(x+27,389,title,36,C['ink'],True,w=395)
    s.text(x+27,442,sub,24,C['ink'],w=395)
    for i in range(3):
        s.rect(x+30+i*115,524,92,80,C[col],r=12)
        s.line([(x+48+i*115,548),(x+101+i*115,548)],C['ink'],3)
        s.line([(x+48+i*115,568),(x+88+i*115,568)],C['ink'],3)
s.arrow(515,468,575,468,C['ink'],4)
s.arrow(1024,468,1084,468,C['ink'],4)
s.pill(408,688,'git add',C['cyan'],w=220)
s.pill(910,688,'git commit',C['lime'],w=240)
s.callout('Erst den Diff lesen. Dann gezielt stagen. Dann committen.',y=759)

# 05 — staging trap
s=Slide('Der Commit nimmt die Auswahl.', 'Aha-Moment / Staging', 'Beispiel: headline.md enthält zuerst Version 1. git add übernimmt Version 1 in den Index. Danach ändern wir im Editor weiter auf Version 2, ohne erneut zu stagen. Der nächste Commit enthält Version 1. Version 2 bleibt als ungestagte Änderung im Working Tree. Frage zuerst, dann zeige mit dem Diagramm die Antwort.',True,subtitle='Nach git add weiterbearbeitet? Diese Änderung ist noch nicht ausgewählt.')
for x,title,version,color in [(100,'EDITOR','Version 2','cyan'),(640,'STAGING','Version 1','purple'),(1160,'COMMIT','Version 1','lime')]:
    s.text(x+135,309,title,24,C['muted'],True,w=300,anchor='middle')
    s.rect(x,367,290,224,C[color],r=22)
    s.text(x+25,406,'headline.md',24,C['ink'],True,w=260,mono=True)
    s.text(x+25,482,version,40,C['ink'],True,w=260)
s.arrow(406,487,619,487,C['muted'],4)
s.text(514,548,'vorher: git add',23,C['muted'],w=235,anchor='middle',mono=True)
s.arrow(951,487,1138,487,C['lime'],5)
s.text(1047,548,'git commit',24,C['lime'],w=210,anchor='middle',mono=True)
s.callout('Nochmals git add ausführen, wenn Version 2 ebenfalls in den Commit soll.')

# 06 — commit snapshot
s=Slide('Ein Commit hält einen Stand fest.', 'Prinzip 04 / Commit', 'Ein Commit verweist auf den kompletten versionierten Projektstand, nicht nur auf den angezeigten Diff. Git speichert identische Objekte effizient, das Bild ist also keine Aussage über vollständige physische Kopien pro Commit. Zum Commit gehören Elternverweise und Metadaten. Die Beispiel-ID ist illustrativ.',subtitle='Der Diff zeigt die Änderung. Der Commit verweist auf den Projektstand.')
s.rect(94,286,642,427,C['white'],r=24)
s.pill(123,313,'COMMIT C',C['lime'],w=210)
for i,name in enumerate(['index.md','program.md','snacks.md']):
    y=398+i*86
    s.rect(123,y,580,66,C['purple'] if i==1 else C['soft'],r=10)
    s.text(147,y+17,name,27,C['ink'],True,w=300,mono=True)
    s.text(677,y+18,'geändert' if i==1 else 'unverändert',23,C['ink'],anchor='end',w=220)
s.line([(780,494),(896,494)],C['ink'],5)
s.node(895,494,'C',C['lime'],35)
s.text(985,322,'ID',23,C['ink'],True,w=180)
s.text(985,356,'a1b2c3d',34,C['ink'],True,w=400,mono=True)
s.text(985,435,'ELTERNCOMMIT',23,C['ink'],True,w=350)
s.text(985,470,'B',36,C['ink'],True,w=300,mono=True)
s.text(985,549,'NACHRICHT',23,C['ink'],True,w=350)
s.text(985,583,'docs: clarify lunch venue',27,C['ink'],True,w=560,mono=True)
s.callout('Ein guter Commit ist ein verständlicher Schritt mit einer hilfreichen Nachricht.')

# 07 branch pointer
s=Slide('Ein Branch ist ein Namensschild.', 'Prinzip 05 / Branch & HEAD', 'Ein Branch ist ein beweglicher Verweis auf einen Commit. Wenn wir einen Branch anlegen, müssen wir nicht das ganze Projekt kopieren. Hier zeigen main und feature auf denselben Commit B. HEAD verweist normalerweise auf den aktuell ausgecheckten Branch, hier feature. Die Linien zwischen Commits zeigen Abstammung; wir zeichnen ältere Stände links.',True,subtitle='Zwei Branch-Namen können auf denselben Commit zeigen.')
s.line([(260,497),(760,497)],C['muted'],8)
s.node(260,497,'A',C['muted'],42)
s.node(760,497,'B',C['lime'],49)
s.tag(760,349,'main',C['lime'],w=170)
s.line([(760,422),(760,448)],C['lime'],4)
s.pill(646,601,'feature',C['cyan'],w=228)
s.line([(760,600),(760,548)],C['cyan'],4)
s.pill(1100,601,'HEAD',C['purple'],w=155)
s.arrow(1080,625,890,625,C['purple'],5)
s.text(1065,422,'Du arbeitest\nauf feature.',35,C['paper'],True,w=370)
s.callout('git switch -c feature  →  neuer Name am aktuellen Commit')

# 08 feature advances
s=Slide('Nur dein Branch bewegt sich.', 'Prinzip 05 / Ein Commit später', 'Vergleiche mit der vorherigen Folie: Ein neuer Commit C wird auf feature erzeugt. main bleibt auf B. HEAD bleibt an feature gebunden und zeigt dadurch indirekt auf C. Ein Commit auf einem Feature-Branch ändert nicht automatisch main. Fragen lassen: Wo steht main jetzt? Wo steht HEAD?',True,subtitle='Ein neuer Commit auf feature lässt main unverändert.')
s.line([(240,498),(735,498),(1230,498)],C['muted'],8)
for x,label,col in [(240,'A','muted'),(735,'B','lime'),(1230,'C','cyan')]: s.node(x,498,label,C[col],43)
s.tag(735,350,'main',C['lime'],w=170)
s.line([(735,422),(735,451)],C['lime'],4)
s.tag(1230,350,'feature',C['cyan'],w=220)
s.line([(1230,422),(1230,451)],C['cyan'],4)
s.pill(1100,610,'HEAD',C['purple'],w=170)
s.line([(1290,634),(1460,634),(1460,374)],C['purple'],4)
s.arrow(1460,374,1355,374,C['purple'],4)
s.text(218,614,'vorhandene History',28,C['muted'],w=550)
s.callout('Der Branch folgt deinen neuen Commits. Andere Branches bleiben stehen.')

# 09 end to end
s=Slide('So wird aus einer Idee Teamarbeit.', 'Der gemeinsame Workflow', 'Führe den gesamten Weg einmal vor: Issue beschreibt die Aufgabe, Branch grenzt die Arbeit ab, Commits machen Schritte nachvollziehbar, Push veröffentlicht. Im Pull Request folgen Review und Nachbesserung. Erst nach finalem Review wird integriert. Ein neuer Push auf denselben Branch aktualisiert den vorhandenen PR.',subtitle='Kleine Änderungen. Sichtbarer Kontext. Eine bewusste Freigabe.')
items=[('01','Issue','Ziel klären','soft'),('02','Branch','Arbeit starten','cyan'),('03','Commits','Schritte festhalten','purple'),('04','Push','Stand teilen','cyan'),('05','PR','Änderung erklären','lime'),('06','Review','prüfen & verbessern','coral'),('07','Merge','integrieren','lime')]
for i,(n,title,sub,col) in enumerate(items):
    x=72+i*210
    s.rect(x,339,193,247,C[col],r=20)
    s.text(x+21,361,n,22,C['ink'],True,w=155)
    s.text(x+21,433,title,29,C['ink'],True,w=175)
    s.text(x+21,501,sub.replace(' & ',' &\n') if i==5 else sub,20,C['ink'],w=175)
    if i<6:s.arrow(x+196,463,x+207,463,C['ink'],2)
s.line([(1300,615),(1300,675),(595,675),(595,615)],C['ink'],4)
s.text(953,690,'Feedback → neuer Commit → gleicher PR',26,C['ink'],w=700,anchor='middle')
s.callout('main bleibt der gemeinsam geprüfte Stand.')

# 10 exercise 1
s=Slide('Euer erster Pull Request.', 'Praxis 01 / Minute 15–35', 'Lasse diese Folie während Runde 1 stehen. Ein PR pro Zweierteam. Team 1 bearbeitet index.md, 2 program.md, 3 snacks.md, 4 team.md, 5 status-board.md. Die genaue Mission steht in exercises/01-warmup.md. Nach dem ersten Commit wechselt das Keyboard. Mindestens zwei sinnvolle Commits, PR mit Ziel und Prüfschritten. Noch nicht mergen.',True,subtitle='Ein Auftrag pro Team. Nach dem ersten Commit wechselt das Keyboard.',minutes=20)
s.text(77,283,'20',169,C['lime'],True,w=340)
s.text(82,456,'MINUTEN',31,C['lime'],True,w=320)
for i,(name,col) in enumerate([('Startseite','cyan'),('Programm','lime'),('Snacks','purple'),('Kontakte','coral'),('Status','cyan')]):
    y=282+i*86
    s.circle(500,y+27,26,C[col]); s.text(500,y+8,str(i+1),26,C['ink'],True,w=50,anchor='middle')
    s.text(548,y+8,name,29,C['paper'],True,w=360)
s.rect(963,282,565,410,C['panel'],r=26)
for i,(n,txt) in enumerate([('1','Branch erstellen'),('2','Zwei sinnvolle Commits'),('3','Push + PR beschreiben')]):
    s.text(998,330+i*109,n,27,C['lime'],True,w=50)
    s.text(1051,330+i*109,txt,30,C['paper'],True,w=435)
s.callout('Checkpoint: 5 offene PRs. Ziel und Prüfung sind beschrieben. Noch nicht mergen.')

# 11 PR conversation
s=Slide('Ein PR macht Entscheidungen sichtbar.', 'Prinzip 06 / Review', 'Ein hilfreiches Review begründet seine Rückfrage anhand der Wirkung für Nutzer:innen. Zeige die fehlende Ortsangabe und den Kommentar. Die Autor:innen ergänzen den Ort und pushen auf denselben Branch. Der PR aktualisiert sich. Reviewer prüfen den neuen Stand. Comment ist Rückmeldung; Request changes fordert Nachbesserung; Approve gibt den geprüften Stand frei.',subtitle='Der Diff zeigt das Was. Die Diskussion klärt das Warum.')
s.rect(83,280,829,439,C['white'],r=24)
s.pill(114,310,'program.md',C['soft'],w=246)
s.rect(113,400,769,93,'#FADDD5',r=12)
s.text(134,430,'− 16:30 Filmwerkstatt',31,C['ink'],True,w=730,mono=True)
s.rect(113,514,769,120,'#E5F4BF',r=12)
s.text(134,550,'+ 16:30 Filmwerkstatt, Raum B',29,C['ink'],True,w=730,mono=True)
s.rect(981,327,528,211,C['purple'],r=23)
s.text(1010,355,'REVIEW',21,C['ink'],True,w=435)
s.text(1010,409,'„Wo findet der neue Slot statt?\nBitte den Ort ergänzen.“',28,C['ink'],True,w=455)
s.pill(981,584,'prüfen',C['soft'],w=146)
s.arrow(1145,608,1200,608,C['ink'],4)
s.pill(1225,584,'Approve',C['lime'],w=200)
s.callout('Rückmeldung → Nachbesserung → erneute Prüfung → Freigabe')

# 12 review pause
s=Slide('Jetzt prüft das nächste Team.', 'Praxis 02 / Minute 35–55', 'Lasse diese Folie während Runde 2 stehen. Review-Ring 1 zu 2, 2 zu 3, 3 zu 4, 4 zu 5, 5 zu 1. Jede Person schreibt mit dem eigenen Konto einen begründeten Beitrag. Pro Team eine konkrete Stärke und eine Frage oder Verbesserung. Entwürfe absenden. Autor:innen übernehmen eine sinnvolle Verbesserung, Reviewer prüfen erneut. Ab Minute 50 nacheinander Create a merge commit verwenden, danach lokalen main aktualisieren.',True,subtitle='Beide Personen schreiben einen sichtbaren Review-Beitrag.',minutes=20)
center=(430,480); pts=[]
for i in range(5):
    a=-math.pi/2+i*2*math.pi/5;pts.append((center[0]+220*math.cos(a),center[1]+195*math.sin(a)))
for i,p in enumerate(pts):
    q=pts[(i+1)%5];dx=q[0]-p[0];dy=q[1]-p[1];d=math.hypot(dx,dy)
    s.arrow(p[0]+dx/d*45,p[1]+dy/d*45,q[0]-dx/d*48,q[1]-dy/d*48,C['muted'],5)
for i,(x,y) in enumerate(pts):s.node(x,y,str(i+1),C['cyan'] if i%2 else C['lime'],40)
s.text(430,453,'REVIEW-\nRING',34,C['paper'],True,w=230,anchor='middle')
for i,(title,sub,col) in enumerate([('1 Stärke','konkret begründen','cyan'),('1 Verbesserung','Wirkung erklären','purple'),('Finales Review','nachbessern → prüfen → mergen','lime')]):
    x=835;y=282+i*150
    s.rect(x,y,663,130,C['panel'],r=19)
    s.text(x+28,y+24,title,33,C[col],True,w=610)
    s.text(x+28,y+75,sub,27,C['paper'],w=610)
s.callout('Entwürfe absenden. Erst nach finalem Review: Create a merge commit.')

# 13 fetch pull
s=Slide('Holen ist noch nicht integrieren.', 'Prinzip 07 / Fetch & Pull', 'Die drei Zeilen gehören zu verschiedenen Orten: Server-Branch main, lokale Remote-Tracking-Ref origin/main und lokaler Branch main. Nach Fetch kann origin/main auf C stehen, während lokaler main auf B bleibt. Pull führt zunächst Fetch und dann einen Integrationsversuch aus. Im Workshop pull --ff-only: bei Divergenz bricht der Integrationsschritt ab. Die Folie zeigt den Zustand nach einem erfolgreichen Fetch und vor dem Pull.',subtitle='Zustand nach git fetch origin: Dein Arbeitsbranch kann noch zurückliegen.')
rows=[(322,'Server: main',C['cyan'],True),(465,'Lokal: origin/main',C['purple'],True),(608,'Lokal: main',C['lime'],False)]
for y,label,col,to_c in rows:
    s.text(84,y-19,label,29,C['ink'],True,w=490,mono=True)
    s.line([(657,y),(1330 if to_c else 1000,y)],C['ink'],5)
    for x,l in [(670,'A'),(1000,'B')]+([(1330,'C')] if to_c else []):s.node(x,y,l,col,30)
    if not to_c:
        s.line([(1035,y),(1315,y)],'#B7C5BF',4,True)
        s.text(1240,658,'noch nicht integriert',23,C['ink'],w=340,anchor='middle')
s.callout('fetch holt den Stand. pull holt und versucht zusätzlich zu integrieren.')

# 14 conflict beginning
s=Slide('Gleicher Start. Zwei Entscheidungen.', 'Prinzip 08 / Konflikt', 'Ab Minute 55 erzeugt jede Person mit bash scripts/create-labs.sh eigene Labore und öffnet conflict. Das Skript hat zwei Branches vom selben Ausgangspunkt erstellt. Team A verlegt den Lunch in die Mensa und ist bereits im lokalen main. Team B ändert Zeit und Angebot. Beide haben dieselbe Zeile verschieden geändert. Gleiche Datei allein garantiert keinen Konflikt; hier ist der Konflikt gezielt konstruiert.',subtitle='Beide Branches haben dieselbe Zeile unterschiedlich verändert.')
s.rect(85,405,425,146,C['soft'],r=20)
s.text(110,432,'GEMEINSAME BASIS',22,C['ink'],True,w=370)
s.text(110,478,'12:00 · Innenhof',34,C['ink'],True,w=370)
s.line([(510,478),(640,478),(760,338),(835,338)],C['ink'],5)
s.line([(640,478),(760,612),(835,612)],C['ink'],5)
s.rect(835,261,680,163,C['lime'],r=22)
s.text(865,289,'TEAM A → MAIN',22,C['ink'],True,w=600)
s.text(865,337,'12:00 · Mensa',37,C['ink'],True,w=610)
s.rect(835,531,680,184,C['cyan'],r=22)
s.text(865,558,'TEAM B',22,C['ink'],True,w=600)
s.text(865,601,'12:30 · Innenhof\nvegetarisches Buffet',34,C['ink'],True,w=610)
s.callout('Git erkennt den Textkonflikt. Ihr klärt die gültige Information für die Gäste.')

# 15 conflict markers
s=Slide('Git stoppt. Ihr entscheidet.', 'Prinzip 08 / Konflikt lesen', 'Führt im frischen conflict-Labor auf team-b git merge main aus. Der erwartete Konflikt stoppt die automatische Zusammenführung. Oben steht in diesem Merge HEAD beziehungsweise team-b, unten der eingehende main. Beim Rebase kann die ours/theirs-Zuordnung anders wirken; nicht verallgemeinern. Die Festivalleitung bestätigt Mensa, 12:30 Uhr und vegetarisches Buffet. Die endgültige Zeile noch nicht verraten.',True,subtitle='Im lokalen Konfliktlabor, auf team-b: git merge main')
s.rect(73,275,1070,411,C['panel'],r=22)
code=[('<<<<<<< HEAD','coral'),('- 12:30 Lunch im Innenhof mit vegetarischem Buffet','cyan'),('=======','coral'),('- 12:00 Lunch in der Mensa','lime'),('>>>>>>> main','coral')]
for i,(line,col) in enumerate(code):s.text(105,309+i*73,line,30,C[col],True,w=1000,mono=True)
s.text(1190,300,'VORGABEN',23,C['muted'],True,w=325)
for y,label,col in [(357,'12:30','cyan'),(452,'Mensa','lime'),(547,'vegetarisch','purple')]:s.pill(1180,y,label,C[col],w=328,size=30)
s.callout('Die Marker zeigen Varianten. Die fachlich richtige Endfassung bestimmt ihr.')

# 16 conflict exercise
s=Slide('Baut die gültige Lunch-Zeile.', 'Praxis 03 / Minute 55–80', 'Konfliktrunde insgesamt 25 Minuten, inklusive der vorherigen Einführung: fünf Minuten Labore und Graph, fünf Minuten Konflikt lesen, sieben Minuten lösen, acht Minuten erklären beziehungsweise zweiten Durchlauf am gemeinsamen Laptop durchführen. Pro Person ein eigener Durchlauf. Nach Bearbeitung git diff --check, git add program.md, git diff --staged und Merge committen. Erfolg: alle drei Vorgaben erfüllt, keine Marker, sauberer Status und Merge-Commit mit zwei Eltern.',True,subtitle='Jede Person löst einen echten Konflikt. Danach erklärt ihr euch die Entscheidung.',minutes=25)
s.text(77,295,'25',167,C['coral'],True,w=370)
s.text(83,469,'MINUTEN GESAMT',26,C['coral'],True,w=350)
steps=[('01','Lesen','Status + beide Varianten','cyan'),('02','Entscheiden','Zeit, Ort und Angebot verbinden','purple'),('03','Abschliessen','prüfen → add → commit','lime')]
for i,(n,title,sub,col) in enumerate(steps):
    y=282+i*149
    s.rect(534,y,977,132,C['panel'],r=20)
    s.pill(560,y+39,n,C[col],w=86)
    s.text(680,y+21,title,34,C[col],True,w=760)
    s.text(680,y+74,sub,28,C['paper'],w=760)
s.callout('Fertig: keine Marker, alle drei Vorgaben erfüllt, sauberer git status.')

# 17 resolution
s=Slide('Drei Anforderungen. Eine Lösung.', 'Auflösung / Erst nach der Praxis', 'Erst nach dem eigenen Lösungsversuch zeigen. Die gemeinsame Fassung lautet 12:30 Lunch in der Mensa mit vegetarischem Buffet. Das Zusammensetzen ist eine inhaltliche Entscheidung. git add markiert als gelöst, prüft aber keine fachliche Richtigkeit. Im Graph zeigt main weiter auf A, team-b nach dem Merge auf M; M hat A und B als Eltern. Das lokale Beispiel braucht keinen Push.',subtitle='Eine gültige Endfassung enthält die Informationen aus beiden Änderungen.')
s.rect(73,272,1455,150,C['lime'],r=24)
s.text(108,307,'12:30 · Lunch in der Mensa',43,C['ink'],True,w=1370)
s.text(108,364,'mit vegetarischem Buffet',32,C['ink'],w=1370)
s.line([(218,590),(550,516),(1130,590)],'#6A8630',8)
s.line([(218,590),(550,681),(1130,590)],'#248293',8)
for x,y,l,col in [(218,590,'O','soft'),(550,516,'A','lime'),(550,681,'B','cyan'),(1130,590,'M','purple')]:s.node(x,y,l,C[col],32)
s.pill(487,450,'main',C['lime'],w=130)
s.pill(1206,565,'team-b',C['cyan'],w=220)
s.text(930,687,'M hat zwei Eltern: A und B.',28,C['ink'],w=570)
s.callout('Konfliktfrei heisst: Git konnte zusammenführen. Richtig heisst: Ihr habt geprüft.')

# 18 FF before after
s=Slide('Fast-Forward: Das Schild rückt nach.', 'Prinzip 09 / Fast-Forward', 'Wechsel ins fast-forward-Labor. Zeige vorher und nachher: A ist ein Vorfahr von B. Das Merge muss keinen zusätzlichen Commit erzeugen, sondern kann main auf B setzen. feature/info und main zeigen anschliessend auf denselben bereits vorhandenen Commit. IDs vergleichen. Nicht mit konfliktfrei gleichsetzen: Das ist eine Aussage über die History.',True,subtitle='Die Commits bleiben dieselben. Nur der Zielbranch bewegt sich.')
for y,word in [(365,'VORHER'),(610,'NACHHER')]:
    s.text(81,y-15,word,24,C['muted'],True,w=280)
    s.line([(560,y),(1240,y)],C['muted'],7)
    s.node(560,y,'A',C['muted'],35);s.node(1240,y,'B',C['cyan'],35)
s.pill(500,273,'main',C['lime'],w=130)
s.pill(1115,273,'feature/info',C['cyan'],w=256)
s.pill(1070,673,'main',C['lime'],w=130)
s.pill(1220,673,'feature/info',C['cyan'],w=282)
s.arrow(614,467,1204,467,C['lime'],6)
s.text(902,414,'git merge --ff-only feature/info',28,C['lime'],True,w=710,anchor='middle',mono=True)
s.callout('Kein neuer Commit. main steht jetzt auf dem vorhandenen Commit B.')

# 19 divergence
s=Slide('Beide Linien sind weitergelaufen.', 'Prinzip 10 / Divergenz', 'Im rebase-Labor stehen main und feature auf eigenen Nachfahren derselben Basis. B liegt nicht in der Abstammung von C. Auf main git merge --ff-only feature/volunteers ausführen: erwartete Ablehnung. Es läuft danach kein Merge, daher ist kein merge --abort nötig. Die Branches ändern verschiedene Dateien; die Ablehnung ist kein Dateikonflikt.',True,subtitle='Fast-Forward ist hier unmöglich – auch ohne Dateikonflikt.')
s.line([(260,506),(675,356),(1150,356)],C['lime'],8)
s.line([(260,506),(675,643),(1150,643)],C['cyan'],8)
s.node(260,506,'A',C['muted'],39)
s.node(1150,356,'B',C['lime'],39)
s.node(1150,643,'C',C['cyan'],39)
s.pill(1215,330,'main',C['lime'],w=170)
s.pill(1215,617,'feature',C['cyan'],w=237)
s.text(629,294,'eigener neuer Commit',27,C['lime'],w=510)
s.text(629,684,'eigener neuer Commit',27,C['cyan'],w=510)
s.pill(612,465,'--ff-only: Stopp',C['coral'],w=416,size=30)
s.callout('Divergenz beschreibt die History. Ein Konflikt betrifft die Zusammenführung.')

# 20 rebase
s=Slide('Rebase: gleiche Änderung, neue Basis.', 'Prinzip 11 / Rebase', 'Auf dem Feature-Branch wird git rebase main ausgeführt. C wird auf B erneut angewendet; es entsteht C′ mit anderem Elterncommit und damit anderer ID. main bleibt auf B. Ein lokaler Rebase kann in anderen Ausgangssituationen ohne Änderungen auskommen; hier ändert sich die Basis bewusst. before-rebase hält im Labor den alten Commit C erreichbar, deshalb zeigt git log --all diesen weiterhin.',True,subtitle='Der Feature-Branch wird neu aufgebaut. main bleibt noch stehen.')
s.text(85,276,'VORHER',22,C['muted'],True,w=260)
s.line([(350,439),(650,354),(1070,354)],C['lime'],7)
s.line([(350,439),(650,529),(1070,529)],C['cyan'],7)
for x,y,l,col in [(350,439,'A','muted'),(1070,354,'B','lime'),(1070,529,'C','cyan')]:s.node(x,y,l,C[col],31)
s.pill(1150,330,'main',C['lime'],w=140)
s.pill(1150,505,'feature',C['cyan'],w=215)
s.text(85,657,'NACHHER',22,C['muted'],True,w=260)
s.line([(350,671),(900,671),(1390,671)],C['muted'],7)
for x,l,col in [(350,'A','muted'),(900,'B','lime'),(1390,"C′",'cyan')]:s.node(x,671,l,C[col],34)
s.pill(830,579,'main',C['lime'],w=140)
s.pill(1280,579,'feature',C['cyan'],w=215)
s.text(709,456,'git rebase main',27,C['purple'],True,w=440,anchor='middle',mono=True)
s.callout('C → C′: neue Basis, neuer Elterncommit, neue Commit-ID.')

# 21 then fast-forward
s=Slide('Jetzt kann main folgen.', 'Prinzip 11 / Nach dem Rebase', 'Nach dem Rebase wechseln wir auf main und führen git merge --ff-only feature/volunteers aus. main ist nun Vorfahr des Feature-Commits; Fast-Forward gelingt. Rebase und anschliessende Integration sind zwei getrennte Schritte. Frage: Welcher Schritt hat die neue Commit-ID erzeugt? Der Rebase, nicht das Fast-Forward. Nur aktive Branches dargestellt, der Vergleichsbranch before-rebase bleibt im Labor vorhanden.',True,subtitle='Zuerst Rebase auf feature. Danach Fast-Forward auf main.')
s.line([(258,448),(782,448),(1310,448)],C['muted'],8)
for x,l,col in [(258,'A','muted'),(782,'B','muted'),(1310,"C′",'cyan')]:s.node(x,448,l,C[col],43)
s.text(770,266,'vorher',23,C['muted'],w=180,anchor='middle')
s.pill(680,315,'main',C['muted'],w=180)
s.arrow(890,339,1205,339,C['lime'],6)
s.text(1305,266,'nachher',23,C['lime'],w=180,anchor='middle')
s.pill(1215,315,'main',C['lime'],w=180)
s.pill(1187,535,'feature',C['cyan'],w=245)
s.text(160,620,'1  git switch main',33,C['paper'],True,w=1250,mono=True)
s.text(160,677,'2  git merge --ff-only feature/volunteers',33,C['lime'],True,w=1330,mono=True)
s.callout('Das Fast-Forward erzeugt keinen weiteren Commit.')

# 22 comparison
s=Slide('Drei Wege. Drei Geschichten.', 'Einordnung / Merge · Rebase · Squash', 'Alle drei Skizzen starten gedanklich bei O: main ist bis B gelaufen; der Feature-Zweig hat C und D. Links Merge: M verbindet B und D, beide Linien bleiben sichtbar. Mitte Rebase: C′ und D′ werden auf B neu aufgebaut und danach integriert. Rechts Squash-Integration: ein neuer S fasst den Inhalt des Features auf B zusammen. Gleiche resultierende Dateien sind im konfliktfreien Beispiel möglich, die History unterscheidet sich. GitHub Rebase and merge erzeugt neue IDs und ist nicht identisch mit lokalem --ff-only.',subtitle='Entscheidet nach Zusammenarbeit und gewünschter Nachvollziehbarkeit.')
for x,t,col in [(73,'Merge-Commit','lime'),(569,'Rebase + FF','cyan'),(1065,'Squash','purple')]:
    s.rect(x,284,462,429,C['white'],r=24)
    s.pill(x+24,309,t,C[col],w=414,size=29)
# mini merge
s.line([(116,499),(247,440),(454,499)],'#6A8630',5)
s.line([(116,499),(247,588),(353,588),(454,499)],'#248293',5)
for x,y,l,c in [(116,499,'O','soft'),(247,440,'B','lime'),(247,588,'C','cyan'),(353,588,'D','cyan'),(454,499,'M','purple')]:s.node(x,y,l,C[c],24)
# rebase chain
s.line([(612,508),(990,508)],C['ink'],5)
for x,l,c in [(612,'O','soft'),(738,'B','lime'),(864,"C′",'cyan'),(990,"D′",'cyan')]:s.node(x,508,l,C[c],24)
# squash chain
s.line([(1118,508),(1458,508)],C['ink'],5)
for x,l,c in [(1118,'O','soft'),(1288,'B','lime'),(1458,'S','purple')]:s.node(x,508,l,C[c],25)
s.text(102,655,'zwei Linien verbinden',25,C['ink'],True,w=390)
s.text(598,655,'Änderungen neu aufbauen',25,C['ink'],True,w=425)
s.text(1094,655,'Schritte zusammenfassen',25,C['ink'],True,w=425)
s.callout('Gemeinsam genutzte History nur nach Absprache umschreiben.')

# 23 labs
s=Slide('Lest die History.', 'Praxis 04 / Minute 80–100', 'Diese 20 Minuten umfassen die vorherigen kurzen Erklärungen und das eigene Ausführen, keine zusätzlichen 20 Minuten. Die Labore aus Runde 3 sind bereits erstellt. Für Fast-Forward IDs vergleichen; im Rebase-Labor erwartete FF-Ablehnung, Rebase, neue ID und erfolgreiches FF. Squash nur als Bonus mit fünf bis zehn Minuten zusätzlich. Alle Arbeiten sind lokal ohne Remote; kein Force Push erforderlich.',True,subtitle='In den lokalen Laboren: vorhersagen → ausführen → Graph vergleichen.',minutes=20)
for i,(title,sub,col,n) in enumerate([('Fast-Forward','Kein neuer Commit.','lime','01'),('Rebase','Neue Basis. Neue ID.','cyan','02'),('Squash','Bonus: 3 → 1 Commit.','purple','03')]):
    x=73+i*498
    s.rect(x,301,460,367,C['panel'],r=26)
    s.pill(x+28,329,n,C[col],w=87)
    s.text(x+29,411,title,43,C[col],True,w=405)
    s.text(x+29,484,sub,28,C['paper'],w=405)
    s.line([(x+56,594),(x+393,594)],C[col],5)
    for dx in [65,229,393]:s.circle(x+dx,594,14,C[col])
s.callout('Ohne Remote. Kein Push nötig. Squash nur bei zusätzlicher Zeit.')

# 24 quiz
s=Slide('Was passiert als Nächstes?', 'Lerncheck / Minute 100–110', 'Erst 30 Sekunden still nachdenken, dann im Paar besprechen. Antworten: 1 Nur der gestagte Stand kommt in den Commit; danach bearbeitete ungestagte Änderungen bleiben im Working Tree. 2 Nach Fetch ist origin/main aktualisiert, lokaler main aber noch nicht notwendigerweise integriert. 3 Nein, ein konfliktfreier Merge beweist keine fachliche Richtigkeit. Zusatzfrage: Wann ist FF möglich? Wenn der Zielstand Vorfahr des einzubindenden Commits ist.',subtitle='Erst selbst entscheiden. Dann im Paar begründen.')
for i,(n,title,desc,col) in enumerate([('01','add → edit → commit','Welcher Stand wird gespeichert?','purple'),('02','fetch ist fertig','Ist main damit schon aktuell?','cyan'),('03','Merge ohne Konflikt','Ist das Ergebnis deshalb richtig?','lime')]):
    y=272+i*153
    s.rect(73,y,1455,128,C['white'],r=22)
    s.pill(100,y+37,n,C[col],w=90)
    s.text(221,y+24,title,34,C['ink'],True,w=650,mono=i<2)
    s.text(221,y+76,desc,28,C['ink'],w=1150)
s.callout('Eine gute Erklärung beschreibt den Zustand vor und nach dem Befehl.')

# 25 ending
s=Slide('Erst verstehen. Dann integrieren.', 'Transfer / Minute 110–120', 'Sammelt drei konkrete Regeln für eure nächste Zusammenarbeit. Lasst die Gruppe formulieren: Status und Branch zuerst lesen; kleine verständliche Änderungen; Reviews begründen und den finalen Stand prüfen. Wer noch keinen eigenen Konflikt gelöst hat, plant genau diesen nächsten Übungsschritt. Verweise auf Glossar, Spickzettel und Drehbuch im Repository. Letzte fünf Minuten sind Puffer.',True,subtitle='Welche drei Regeln nehmt ihr in eure nächste Zusammenarbeit mit?')
for i,(big,title,sub,col) in enumerate([('01','Zustand lesen','Repository · Branch · Status','cyan'),('02','Kleine Schritte','klare Commits · verständliche PRs','purple'),('03','Gemeinsam prüfen','Feedback · Nachbesserung · Freigabe','lime')]):
    x=73+i*500
    s.text(x,280,big,95,C[col],True,w=450)
    s.text(x,422,title,38,C['paper'],True,w=462)
    s.text(x,488,sub.replace(' · ','\n'),27,C['muted'],w=440)
s.rect(73,704,1455,100,C['lime'],r=20)
s.text(106,728,'github.com/Witzelfitz/git-workflow-dojo',37,C['ink'],True,w=1380,mono=True)

# Render SVGs with the same scene graph used by PowerPoint.
def svg(slide):
    out=[f'<svg xmlns="http://www.w3.org/2000/svg" viewBox="0 0 {W} {H}" width="{W}" height="{H}" role="img" aria-label="{html.escape(slide.title,quote=True)}">']
    for d in slide.shapes:
        if d['k']=='rect':
            out.append(f'<rect x="{d["x"]}" y="{d["y"]}" width="{d["w"]}" height="{d["h"]}" rx="{d["r"]}" fill="{d["fill"]}" stroke="{d["stroke"] or "none"}" stroke-width="{d["sw"]}"/>')
        elif d['k']=='circle':
            out.append(f'<circle cx="{d["x"]}" cy="{d["y"]}" r="{d["r"]}" fill="{d["fill"]}" stroke="{d["stroke"] or "none"}" stroke-width="{d["sw"]}"/>')
        elif d['k']=='line':
            pts=' '.join(f'{x},{y}' for x,y in d['points'])
            dash=' stroke-dasharray="11 10"' if d['dash'] else ''
            out.append(f'<polyline points="{pts}" fill="none" stroke="{d["color"]}" stroke-width="{d["width"]}" stroke-linecap="round" stroke-linejoin="round"{dash}/>')
        else:
            family='Courier New, monospace' if d['mono'] else 'Arial, Helvetica, sans-serif'
            txts=d['text'].split('\n')
            out.append(f'<text x="{d["x"]}" y="{d["y"]+d["size"]*.86}" font-family="{family}" font-size="{d["size"]}" font-weight="{700 if d["bold"] else 400}" fill="{d["color"]}" text-anchor="{d["anchor"]}" data-max-width="{d["w"]}">')
            for j,t in enumerate(txts):out.append(f'<tspan x="{d["x"]}" dy="{0 if j==0 else d["size"]*1.23}">{html.escape(t)}</tspan>')
            out.append('</text>')
    return ''.join(out)+'</svg>'

pages=[]
for n,s in enumerate(slides,1):
    pages.append(f'<section class="slide" id="slide-{n}" data-document-role="page" data-label="{html.escape(s.title,quote=True)}" data-speaker-notes="{html.escape(s.notes,quote=True)}">{svg(s)}</section>')
base_css='''*{box-sizing:border-box}html,body{margin:0;background:#10232D;font-family:Arial,Helvetica,sans-serif}.slide{width:1600px;height:900px;position:relative;overflow:hidden;background:#F5F3EA}.slide svg{display:block;width:100%;height:100%}@page{size:1600px 900px;margin:0}@media print{.slide{break-after:page;page-break-after:always}.slide:last-child{break-after:auto}nav,#notes,#timer{display:none!important}body{background:white}.slide{display:block!important;transform:none!important;position:relative!important}}'''
static='<!doctype html><html lang="de"><head><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1"><title>Git wird greifbar – Git Workflow Dojo</title><style>'+base_css+'</style></head><body>'+''.join(pages)+'</body></html>'
(OUT/'print.html').write_text(static)

# Canva imports native HTML text boxes; complete SVG slides would flatten text.
def canva_page(slide,number):
    pieces=[]
    for d in slide.shapes:
        if d['k'] in ('rect','circle'):
            if d['k']=='circle':
                x,y,w,h,r=d['x']-d['r'],d['y']-d['r'],2*d['r'],2*d['r'],'50%'
            else:
                x,y,w,h,r=d['x'],d['y'],d['w'],d['h'],f'{d["r"]}px'
            border=f'{d["sw"]}px solid {d["stroke"]}' if d['stroke'] else 'none'
            pieces.append(f'<div style="position:absolute;left:{x}px;top:{y}px;width:{w}px;height:{h}px;background:{d["fill"]};border:{border};border-radius:{r}"></div>')
        elif d['k']=='text':
            x=d['x']-(d['w']/2 if d['anchor']=='middle' else d['w'] if d['anchor']=='end' else 0)
            family='Courier New,monospace' if d['mono'] else 'Arial,Helvetica,sans-serif'
            align={'start':'left','middle':'center','end':'right'}[d['anchor']]
            txt=html.escape(d['text']).replace('\n','<br>')
            pieces.append(f'<p style="position:absolute;left:{x}px;top:{d["y"]-2}px;width:{d["w"]}px;margin:0;font-family:{family};font-size:{d["size"]}px;line-height:1.23;font-weight:{700 if d["bold"] else 400};color:{d["color"]};text-align:{align};white-space:nowrap">{txt}</p>')
        else:
            padding=d['width']+2
            x=min(a for a,b in d['points'])-padding;y=min(b for a,b in d['points'])-padding
            w=max(a for a,b in d['points'])-x+padding;h=max(b for a,b in d['points'])-y+padding
            pts=' '.join(f'{a-x},{b-y}' for a,b in d['points'])
            dash='stroke-dasharray="11 10"' if d['dash'] else ''
            pieces.append(f'<svg xmlns="http://www.w3.org/2000/svg" width="{w}" height="{h}" viewBox="0 0 {w} {h}" style="position:absolute;left:{x}px;top:{y}px;width:{w}px;height:{h}px"><polyline points="{pts}" fill="none" stroke="{d["color"]}" stroke-width="{d["width"]}" stroke-linecap="round" stroke-linejoin="round" {dash}/></svg>')
    return f'<section class="slide" data-document-role="page" data-label="{html.escape(slide.title,quote=True)}" data-speaker-notes="{html.escape(slide.notes,quote=True)}">'+''.join(pieces)+'</section>'

(OUT/'canva-import.html').write_text('<!doctype html><html lang="de"><head><meta charset="utf-8"><title>Git wird greifbar</title><style>'+base_css+'</style></head><body>'+''.join(canva_page(s,i) for i,s in enumerate(slides,1))+'</body></html>')
interaction_css='''@media screen{body{height:100vh;overflow:hidden}.slide{display:none;position:absolute;left:50%;top:calc((100vh - 80px)/2);transform:translate(-50%,-50%) scale(var(--scale,1));transform-origin:center}.slide.active{display:block}nav{position:fixed;bottom:12px;left:50%;transform:translateX(-50%);display:flex;align-items:center;gap:5px;background:#07141eed;border:1px solid #42606e;border-radius:40px;padding:6px 10px;color:#F5F3EA;opacity:.25;transition:opacity .2s}nav:hover,nav:focus-within{opacity:1}button,select{border:0;border-radius:16px;background:#213944;color:#F5F3EA;padding:9px 14px;font-size:13px;cursor:pointer}button:focus-visible,select:focus-visible{outline:3px solid #DAF975}#counter{font-size:13px;padding:0 7px;min-width:58px}#notes{display:none;position:fixed;right:24px;bottom:78px;width:min(560px,90vw);padding:24px;border:1px solid #91a0a6;background:#F5F3EA;color:#172C36;border-radius:18px;font-size:18px;line-height:1.5;max-height:60vh;overflow:auto;box-shadow:0 8px 40px #0007}#notes.open{display:block}#notes h2{font-size:19px;margin:0 0 12px}#notes p{margin:0}#timer{display:none;align-items:center;gap:4px;padding-left:8px;color:#DAF975;font-variant-numeric:tabular-nums}#timer.visible{display:flex}#timer output{font-size:18px;padding-right:6px}nav select{max-width:250px}@media(max-width:1100px){nav select{display:none}}}'''
controls='<nav aria-label="Präsentationssteuerung"><button id="prev" title="Vorherige Folie (←)">←</button><span id="counter"></span><button id="next" title="Nächste Folie (→ oder Leertaste)">→</button><select id="jump" aria-label="Folie auswählen"></select><button id="noteBtn" title="Notizen (N) – auf dem geteilten Bildschirm sichtbar">Notizen · N</button><button id="fullBtn" title="Vollbild (F)">Vollbild · F</button><div id="timer"><output>20:00</output><button id="timerToggle">Start</button><button id="timerReset">↺</button></div></nav><aside id="notes" aria-label="Sprechernotizen"><h2></h2><p></p></aside>'
def phase(i):
    if i < 9: return ('Einstieg',15)
    if i == 9: return ('PR-Runde',20)
    if i < 12: return ('Review-Runde',20)
    if i < 17: return ('Konflikt-Runde',25)
    if i < 23: return ('History-Runde',20)
    return ('Debrief',20)
data=json.dumps([dict(title=s.title,notes=s.notes,phase=phase(i)[0],minutes=phase(i)[1]) for i,s in enumerate(slides)],ensure_ascii=False).replace('</','<\\/')
script=r'''
const data=__DATA__;const slides=[...document.querySelectorAll('.slide')];
const notes=document.getElementById('notes'),jump=document.getElementById('jump');let index=0,running=false,remaining=0,endAt=0,timerPhase=null;
data.forEach((s,i)=>{let o=document.createElement('option');o.value=i;o.textContent=String(i+1).padStart(2,'0')+' · '+s.title;jump.append(o)});
function resize(){document.documentElement.style.setProperty('--scale',Math.min(innerWidth/1600,(innerHeight-80)/900));}
function show(i){index=Math.max(0,Math.min(slides.length-1,i));slides.forEach((s,n)=>s.classList.toggle('active',n===index));document.getElementById('counter').textContent=(index+1)+' / '+slides.length;jump.value=index;notes.querySelector('h2').textContent=data[index].title;notes.querySelector('p').textContent=data[index].notes;history.replaceState(null,'','#'+(index+1));if(timerPhase!==data[index].phase){timerPhase=data[index].phase;running=false;remaining=data[index].minutes*60;}document.getElementById('timer').classList.add('visible');document.getElementById('timer').title=timerPhase+' · Gesamtzeit der Runde';updateTimer();}
function updateTimer(){if(running){remaining=Math.max(0,Math.ceil((endAt-Date.now())/1000));if(!remaining)running=false;}document.querySelector('#timer output').textContent=String(Math.floor(remaining/60)).padStart(2,'0')+':'+String(remaining%60).padStart(2,'0');document.getElementById('timerToggle').textContent=running?'Pause':'Start';}
function toggleNotes(){notes.classList.toggle('open');}
async function fullscreen(){try{if(document.fullscreenElement)await document.exitFullscreen();else await document.documentElement.requestFullscreen();}catch(e){console.warn('Vollbild nicht verfügbar',e)}}
document.getElementById('prev').onclick=()=>show(index-1);document.getElementById('next').onclick=()=>show(index+1);jump.onchange=()=>show(Number(jump.value));document.getElementById('noteBtn').onclick=toggleNotes;document.getElementById('fullBtn').onclick=fullscreen;
document.getElementById('timerToggle').onclick=()=>{if(running){remaining=Math.max(0,Math.ceil((endAt-Date.now())/1000));running=false;}else{if(!remaining)remaining=data[index].minutes*60;endAt=Date.now()+remaining*1000;running=true;}updateTimer();};document.getElementById('timerReset').onclick=()=>{running=false;remaining=data[index].minutes*60;updateTimer();};
addEventListener('keydown',e=>{if(['SELECT','INPUT','TEXTAREA'].includes(e.target.tagName)||(e.target.tagName==='BUTTON'&&e.key===' '))return;if(['ArrowRight','PageDown',' '].includes(e.key)){e.preventDefault();show(index+1)}else if(['ArrowLeft','PageUp'].includes(e.key)){e.preventDefault();show(index-1)}else if(e.key==='Home')show(0);else if(e.key==='End')show(slides.length-1);else if(e.key.toLowerCase()==='n')toggleNotes();else if(e.key.toLowerCase()==='f')fullscreen();});addEventListener('resize',resize);addEventListener('hashchange',()=>show((parseInt(location.hash.slice(1))||1)-1));setInterval(updateTimer,250);resize();show((parseInt(location.hash.slice(1))||1)-1);
'''.replace('__DATA__',data)
# The interactive player has no import annotations; the static Canva file above does.
interactive_pages=[p.replace(' data-document-role="page"','') for p in pages]
player='<!doctype html><html lang="de"><head><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1"><title>Git wird greifbar – Präsentation</title><style>'+base_css+interaction_css+'</style></head><body>'+''.join(interactive_pages)+controls+'<script>'+script+'</script></body></html>'
(OUT/'git-dojo.html').write_text(player)
notes_markdown=('# Git wird greifbar – Sprechernotizen\n\n25 Folien, passend zum 120-Minuten-Dojo. Zeitangaben der Praxisfolien umfassen die vorherigen Kurzimpulse der jeweiligen Runde.\n\n'+''.join(f'## {i:02} · {s.title}\n\n{s.notes}\n\n' for i,s in enumerate(slides,1)))
(OUT/'speaker-notes.md').write_text(notes_markdown.rstrip()+'\n')

# Native editable PowerPoint shapes; no full-slide screenshots.
def make_pptx():
    from pptx import Presentation
    from pptx.util import Inches, Pt
    from pptx.dml.color import RGBColor
    from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
    from pptx.enum.text import PP_ALIGN, MSO_AUTO_SIZE
    from pptx.enum.dml import MSO_LINE_DASH_STYLE
    prs=Presentation();prs.slide_width=Inches(13.333333);prs.slide_height=Inches(7.5)
    px=lambda n: Inches(n/120)
    rgb=lambda h: RGBColor.from_string(h.lstrip('#'))
    for s in slides:
        p=prs.slides.add_slide(prs.slide_layouts[6])
        p.notes_slide.notes_text_frame.text=s.notes
        for d in s.shapes:
            if d['k'] in ['rect','circle']:
                if d['k']=='circle': x,y,w,h=d['x']-d['r'],d['y']-d['r'],2*d['r'],2*d['r'];kind=MSO_SHAPE.OVAL
                else: x,y,w,h=d['x'],d['y'],d['w'],d['h'];kind=MSO_SHAPE.ROUNDED_RECTANGLE if d['r'] else MSO_SHAPE.RECTANGLE
                a=p.shapes.add_shape(kind,px(x),px(y),px(w),px(h));a.fill.solid();a.fill.fore_color.rgb=rgb(d['fill'])
                if d['k']=='rect' and d['r']:
                    a.adjustments[0]=min(.5,d['r']/min(w,h))
                if d['stroke']:a.line.color.rgb=rgb(d['stroke']);a.line.width=Pt(d['sw']*.6)
                else:a.line.fill.background()
            elif d['k']=='line':
                for (x1,y1),(x2,y2) in zip(d['points'],d['points'][1:]):
                    a=p.shapes.add_connector(MSO_CONNECTOR.STRAIGHT,px(x1),px(y1),px(x2),px(y2));a.line.color.rgb=rgb(d['color']);a.line.width=Pt(d['width']*.6)
                    if d['dash']:a.line.dash_style=MSO_LINE_DASH_STYLE.DASH
            else:
                x=d['x']-(d['w']/2 if d['anchor']=='middle' else d['w'] if d['anchor']=='end' else 0)
                lines=d['text'].split('\n');height=len(lines)*d['size']*1.23+10
                a=p.shapes.add_textbox(px(x),px(d['y']-1),px(d['w']),px(height));tf=a.text_frame
                tf.word_wrap=False;tf.auto_size=MSO_AUTO_SIZE.NONE
                tf.margin_left=tf.margin_right=tf.margin_top=tf.margin_bottom=0
                for j,line in enumerate(lines):
                    para=tf.paragraphs[0] if j==0 else tf.add_paragraph();para.text=line
                    para.alignment={'start':PP_ALIGN.LEFT,'middle':PP_ALIGN.CENTER,'end':PP_ALIGN.RIGHT}[d['anchor']]
                    para.space_before=Pt(0);para.space_after=Pt(0);para.line_spacing=1.23
                    para.font.name='Courier New' if d['mono'] else 'Arial';para.font.size=Pt(d['size']*.6);para.font.bold=d['bold'];para.font.color.rgb=rgb(d['color'])
    prs.core_properties.title='Git wird greifbar – Git Workflow Dojo'
    prs.core_properties.subject='Visuelle Workshop-Präsentation: Git-Grundlagen, PR, Review, Konflikte, Fast-Forward und Rebase'
    prs.core_properties.author='Git Workflow Dojo'
    prs.save(OUT/'git-dojo.pptx')
make_pptx()
print(f'Built {len(slides)} slides: HTML player, Canva import, editable PPTX and notes.')
